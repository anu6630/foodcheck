#!/usr/bin/env python3
"""Generate Play Store listing translations via Gemini, then emit CSV/XLSX."""

from __future__ import annotations

import csv
import json
import shutil
import ssl
import time
import urllib.error
import urllib.request
from pathlib import Path

import certifi

SSL_CTX = ssl.create_default_context(cafile=certifi.where())

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
DL = Path.home() / "Downloads"
KEY = (REPO / "apikey_gemini").read_text(encoding="utf-8").strip()
MODEL = "gemini-2.5-flash"

TITLE = "LifeX: Food Ingredients Check"
SHORT = (
    "Scan food labels to spot banned & restricted ingredients—know what’s in a pack."
)
FULL = """LifeX Food helps you understand what’s in packaged food—before you buy or serve it.

Snap a photo of the ingredients label. We read the pack, match substances, and show regulation signals for your market (starting with India)—including banned and restricted additives—not vague “health scores.”

What you can do
• Scan ingredient labels with your camera or photo library
• See clear substance tags with watch signals and regulation context
• Search products and ingredients in the catalog
• Save packs to My list on this device with your stance (Avoid, Occasional, OK, Unsure)
• Optionally tag items for Kids, Adults, Elderly, or Home

Important
LifeX Food is an information tool. It does not provide medical advice. Always follow local regulations and your own judgment for dietary needs.

Know what’s banned, anywhere."""

# Play Console locale codes + natural language name for the model
LOCALES: list[tuple[str, str]] = [
    ("hi-IN", "Hindi (India)"),
    ("bn-BD", "Bengali"),
    ("te-IN", "Telugu"),
    ("mr-IN", "Marathi"),
    ("ta-IN", "Tamil"),
    ("gu-IN", "Gujarati"),
    ("kn-IN", "Kannada"),
    ("ml-IN", "Malayalam"),
    ("pa-IN", "Punjabi (Gurmukhi)"),
    ("or-IN", "Odia"),
    ("as-IN", "Assamese"),
    ("ur", "Urdu"),
    ("de-DE", "German (Germany)"),
    ("fr-FR", "French (France)"),
    ("es-ES", "Spanish (Spain)"),
    ("es-419", "Spanish (Latin America)"),
    ("pt-BR", "Portuguese (Brazil)"),
    ("it-IT", "Italian"),
    ("nl-NL", "Dutch"),
    ("pl-PL", "Polish"),
    ("ru-RU", "Russian"),
    ("tr-TR", "Turkish"),
    ("ar", "Arabic"),
    ("zh-CN", "Simplified Chinese"),
    ("zh-TW", "Traditional Chinese"),
    ("ja-JP", "Japanese"),
    ("ko-KR", "Korean"),
    ("th", "Thai"),
    ("vi", "Vietnamese"),
    ("id", "Indonesian"),
    ("ms-MY", "Malay"),
]


def trim(s: str, n: int) -> str:
    s = s.strip()
    if len(s) <= n:
        return s
    cut = s[: n - 1].rsplit(" ", 1)[0]
    return (cut if cut else s[: n - 1]) + "…"


def gemini_json(prompt: str) -> dict:
    url = (
        f"https://generativelanguage.googleapis.com/v1beta/models/"
        f"{MODEL}:generateContent?key={KEY}"
    )
    body = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {
            "temperature": 0.2,
            "responseMimeType": "application/json",
        },
    }
    data = json.dumps(body).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=data,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    last: Exception | None = None
    for attempt in range(10):
        try:
            with urllib.request.urlopen(req, timeout=120, context=SSL_CTX) as resp:
                raw = json.loads(resp.read().decode("utf-8"))
            text = raw["candidates"][0]["content"]["parts"][0]["text"]
            return json.loads(text)
        except urllib.error.HTTPError as e:
            last = e
            body = e.read().decode("utf-8", errors="replace")[:300]
            wait = 20 * (attempt + 1) if e.code == 429 else 2.0 * (attempt + 1)
            print(f"    retry {attempt+1} http {e.code} sleep {wait:.0f}s {body[:80]!r}")
            time.sleep(wait)
        except (urllib.error.URLError, TimeoutError, KeyError, json.JSONDecodeError) as e:
            last = e
            wait = 3.0 * (attempt + 1)
            print(f"    retry {attempt+1} {type(e).__name__} sleep {wait:.0f}s")
            time.sleep(wait)
    raise RuntimeError(f"Gemini failed: {last}")


def translate_locale(code: str, lang_name: str) -> tuple[str, str, str]:
    prompt = f"""You are a professional Google Play store listing localizer.
Translate the following English Play Store listing into {lang_name}.

Hard rules:
- Keep the brand name "LifeX" and "LifeX Food" untranslated.
- title MUST be <= 30 characters (count Unicode code points / Python len).
- short_description MUST be <= 80 characters.
- full_description MUST be <= 4000 characters and preserve paragraph structure and bullet points (•).
- Natural, native marketing tone for app stores in that locale — not word-for-word.
- Do not add medical claims beyond the source.
- Return ONLY JSON with keys: title, short_description, full_description

SOURCE TITLE:
{TITLE}

SOURCE SHORT DESCRIPTION:
{SHORT}

SOURCE FULL DESCRIPTION:
{FULL}
"""
    obj = gemini_json(prompt)
    title = trim(str(obj["title"]), 30)
    short = trim(str(obj["short_description"]), 80)
    full = str(obj["full_description"]).strip()
    if "LifeX" not in title:
        title = trim("LifeX: " + title, 30)
    if len(full) > 4000:
        full = full[:3999] + "…"
    return title, short, full


def write_outputs(rows: list[tuple[str, str, str, str]], errors: list[tuple[str, str]]) -> None:
    ROOT.mkdir(parents=True, exist_ok=True)

    src = ROOT / "play-store-listing-source-for-ai.csv"
    with src.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, quoting=csv.QUOTE_ALL)
        w.writerow(["key", "text"])
        w.writerow(["title", TITLE])
        w.writerow(["short_description", SHORT])
        w.writerow(["full_description", FULL])

    csv_path = ROOT / "play-store-listing-translations.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, quoting=csv.QUOTE_ALL)
        w.writerow(["language", "title", "short_description", "full_description"])
        w.writerows(rows)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Store listings"
    ws.append(
        [
            "language",
            "title",
            "short_description",
            "full_description",
            "title_len",
            "short_len",
        ]
    )
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, size=11)
    for lang, t, s, f in rows:
        ws.append([lang, t, s, f, len(t), len(s)])
        for cell in ws[ws.max_row]:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 52
    ws.column_dimensions["D"].width = 80
    xlsx = ROOT / "play-store-listing-translations.xlsx"
    wb.save(xlsx)

    fail_txt = (
        "\n".join(f"- {l}: {e}" for l, e in errors) if errors else "None"
    )
    (ROOT / "README.md").write_text(
        f"""# Play Store listing translations

## Files
- `play-store-listing-translations.csv` / `.xlsx` — **{len(rows)}** locales
- `play-store-listing-source-for-ai.csv` — English source for Play Console **Import translations with AI**

## How to use
1. Play Console → Grow → Store presence → Main store listing
2. Click **Import translations with AI** → **Import a file**
3. Upload the CSV or XLSX from this folder (or `~/Downloads`)
4. Review titles (≤30) and short descriptions (≤80) before publishing

## Locales included
Major Indian languages (Hindi, Bengali, Telugu, Marathi, Tamil, Gujarati, Kannada, Malayalam, Punjabi, Odia, Assamese, Urdu) plus markets that strongly prefer local language: German, French, Spanish (ES + LatAm), Portuguese (Brazil), Italian, Dutch, Polish, Russian, Turkish, Arabic, Simplified/Traditional Chinese, Japanese, Korean, Thai, Vietnamese, Indonesian, Malay.

## Generation failures
{fail_txt}
""",
        encoding="utf-8",
    )

    for p in [csv_path, xlsx, src, ROOT / "README.md"]:
        shutil.copy2(p, DL / p.name)
        print("copied", DL / p.name)


def main() -> None:
    rows: list[tuple[str, str, str, str]] = [("en-US", TITLE, SHORT, FULL)]
    errors: list[tuple[str, str]] = []
    cache = ROOT / "translations-cache.json"
    cached: dict[str, dict[str, str]] = {}
    if cache.exists():
        cached = json.loads(cache.read_text(encoding="utf-8"))

    for code, name in LOCALES:
        if code in cached and all(
            k in cached[code] for k in ("title", "short_description", "full_description")
        ):
            t = cached[code]["title"]
            s = cached[code]["short_description"]
            f = cached[code]["full_description"]
            rows.append((code, t, s, f))
            print(f"cache {code} title={len(t)} short={len(s)}")
            continue
        print(f"translating {code} ({name})...")
        try:
            t, s, f = translate_locale(code, name)
            cached[code] = {
                "title": t,
                "short_description": s,
                "full_description": f,
            }
            cache.write_text(
                json.dumps(cached, ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )
            rows.append((code, t, s, f))
            print(f"  ok {code} title={len(t)} short={len(s)}")
            time.sleep(2.0)
        except Exception as e:  # noqa: BLE001
            errors.append((code, str(e)))
            print(f"  FAIL {code}: {e}")
            time.sleep(15)

    write_outputs(rows, errors)
    print(f"done rows={len(rows)} errors={len(errors)}")


if __name__ == "__main__":
    main()
